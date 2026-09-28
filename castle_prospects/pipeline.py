"""Castle Chiropractic prospect pipeline.

Usage (from the repo root):
    python -m castle_prospects.pipeline estimate   # API cost estimate, makes no calls
    python -m castle_prospects.pipeline verify     # Places lookup + drive time for existing records
    python -m castle_prospects.pipeline discover   # find new Tier A organizations
    python -m castle_prospects.pipeline emails     # scrape contact emails from websites
    python -m castle_prospects.pipeline build      # write the workbook + print summary
    python -m castle_prospects.pipeline all        # verify, discover, emails, build

Add --refresh to ignore cached API responses.
"""
import argparse
import csv
import json
import math
import sys
from collections import Counter
from datetime import date

from . import config
from .scoring import classify, clean_name, digits10, name_similarity, priority, review_note, search_query

HISTORY_COLS = [
    "type", "status", "owner", "attempts", "first_contacted", "last_contacted", "last_outcome",
    "last_worked_by", "gatekeeper", "decision_maker", "info_sent_to", "last_note",
    "events_booked", "last_event_date", "active_cadence", "next_step_due",
    "shares_phone_with_another",
]
MATCH_THRESHOLD = 0.6

R_DNC = "Previously asked not to be contacted"
R_CLOSED = "Closed"
R_FAR = "Over 15 min drive"
R_NOFIT = "No audience fit"
R_NOTFOUND = "Not found on Google (needs manual check)"


# --- io ------------------------------------------------------------------------

def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def save_json(name, data):
    config.WORK_DIR.mkdir(parents=True, exist_ok=True)
    (config.WORK_DIR / name).write_text(json.dumps(data, indent=1))


def load_json(name, default=None):
    path = config.WORK_DIR / name
    if not path.exists():
        if default is not None:
            return default
        sys.exit(f"Missing {path} — run the earlier step first.")
    return json.loads(path.read_text())


def split_records():
    """Return (active_records, do_not_contact_records)."""
    all_rows = read_csv(config.ALL_RECORDS_CSV)
    dnc = read_csv(config.DO_NOT_CONTACT_CSV)
    dnc_names = {r["business"].strip() for r in dnc}
    active = [r for r in all_rows if r["business"].strip() not in dnc_names]
    return active, dnc


class DoNotContact:
    def __init__(self, rows):
        self.phones = {digits10(r["phone"]): r for r in rows if digits10(r["phone"])}
        self.names = {clean_name(r["business"]): r for r in rows if clean_name(r["business"])}

    def match(self, name=None, phone=None, place=None):
        for p in (phone, (place or {}).get("nationalPhoneNumber")):
            if digits10(p) in self.phones:
                return self.phones[digits10(p)]
        for n in (name, ((place or {}).get("displayName") or {}).get("text")):
            if n and clean_name(n) in self.names:
                return self.names[clean_name(n)]
        return None


def km_between(a, b):
    lat1, lon1, lat2, lon2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 12742 * math.asin(math.sqrt(h))


def latlng(place):
    loc = place.get("location") or {}
    return (loc.get("latitude"), loc.get("longitude"))


def pretty_phone(phone):
    d = digits10(phone)
    return f"({d[:3]}) {d[3:6]}-{d[6:]}" if d else ""


# --- step: estimate ------------------------------------------------------------

def cmd_estimate(_args):
    active, dnc = split_records()
    per_city = sum(1 for _, pc in config.DISCOVERY_QUERIES if pc)
    single = len(config.DISCOVERY_QUERIES) - per_city
    discover_queries = per_city * len(config.DISCOVERY_CITIES) + single
    # Verification: 1 name search each, plus a phone search for ~25% that don't match.
    verify_calls = round(len(active) * 1.25)
    verify_worst = len(active) * 2
    # Discovery: assume an average of 2 pages per query, 3 worst case.
    discover_calls = discover_queries * 2
    discover_worst = discover_queries * config.DISCOVERY_MAX_PAGES
    # Route elements: one per verified record + one per unique discovered place.
    elements = len(active) + discover_queries * 20
    elements_worst = len(active) + discover_worst * 20

    ts = config.PRICING["text_search_enterprise"]
    rm = config.PRICING["route_matrix_element"]

    def cost(n, sku, free=True):
        billable = max(0, n - sku["free_per_month"]) if free else n
        return billable * sku["per_1000"] / 1000

    text_typ, text_worst = verify_calls + discover_calls + 1, verify_worst + discover_worst + 1
    print(f"Records to verify: {len(active)}  (skipping {len(dnc)} do-not-contact records)")
    print(f"Discovery queries: {discover_queries}\n")
    print(f"{'':32}{'typical':>10}{'worst':>10}")
    print(f"{'Text Search calls (Enterprise)':32}{text_typ:>10}{text_worst:>10}")
    print(f"{'Route matrix elements':32}{elements:>10}{elements_worst:>10}\n")
    def money(x):
        return f"${x:,.2f}".rjust(10)

    print(f"{'Cost at list price':32}{money(cost(text_typ, ts, False) + cost(elements, rm, False))}"
          f"{money(cost(text_worst, ts, False) + cost(elements_worst, rm, False))}")
    print(f"{'Cost after monthly free calls':32}{money(cost(text_typ, ts) + cost(elements, rm))}"
          f"{money(cost(text_worst, ts) + cost(elements_worst, rm))}")
    print("\nRe-runs reuse cached responses in cache/ and cost ~$0. "
          "Prices: confirm at https://mapsplatform.google.com/pricing")


# --- step: verify --------------------------------------------------------------

def best_match(record, places, origin):
    phone = digits10(record["phone"])
    nearby = [p for p in places
              if p.get("location") and km_between(origin, latlng(p)) * 1000 <= config.SEARCH_RADIUS_METERS]
    for p in nearby:
        if phone and digits10(p.get("nationalPhoneNumber")) == phone:
            return p, "phone"
    scored = sorted(((name_similarity(record["business"], (p.get("displayName") or {}).get("text", "")), p)
                     for p in nearby), key=lambda t: -t[0])
    if scored and scored[0][0] >= MATCH_THRESHOLD:
        return scored[0][1], f"name ({scored[0][0]:.2f})"
    return None, ""


def cmd_verify(args):
    from .gapi import GoogleClient
    client = GoogleClient(refresh=args.refresh)
    origin, origin_addr = client.geocode_origin(config.CLINIC_ADDRESS)
    print(f"Clinic located: {origin_addr}")
    active, _ = split_records()
    out = []
    for i, rec in enumerate(active, 1):
        res = client.text_search(search_query(rec["business"]), center=origin,
                                 radius=config.SEARCH_RADIUS_METERS, page_size=5)
        place, method = best_match(rec, res.get("places", []), origin)
        if not place and digits10(rec["phone"]):
            res = client.text_search(pretty_phone(rec["phone"]), center=origin,
                                     radius=config.SEARCH_RADIUS_METERS, page_size=5)
            for p in res.get("places", []):
                if digits10(p.get("nationalPhoneNumber")) == digits10(rec["phone"]):
                    place, method = p, "phone search"
                    break
        out.append({"record": rec, "place": place, "method": method, "drive_minutes": None})
        if i % 25 == 0:
            print(f"  verified {i}/{len(active)}")

    matched = [o for o in out if o["place"]]
    drives = client.drive_minutes(origin, [latlng(o["place"]) for o in matched])
    for o, d in zip(matched, drives):
        o["drive_minutes"] = d
    save_json("verified.json", {"origin": origin, "drive_source": client.drive_source, "rows": out})
    print(f"Matched {len(matched)}/{len(out)} records. API calls: {client.calls}")


# --- step: discover ------------------------------------------------------------

def cmd_discover(args):
    from .gapi import GoogleClient
    client = GoogleClient(refresh=args.refresh)
    verified = load_json("verified.json")
    origin = tuple(verified["origin"])
    all_rows, dnc_rows = read_csv(config.ALL_RECORDS_CSV), read_csv(config.DO_NOT_CONTACT_CSV)
    dnc = DoNotContact(dnc_rows)
    known_ids = {o["place"]["id"] for o in verified["rows"] if o["place"]}
    known_phones = {digits10(r["phone"]) for r in all_rows} | {
        digits10(o["place"].get("nationalPhoneNumber")) for o in verified["rows"] if o["place"]}
    unverified_names = {clean_name(o["record"]["business"]) for o in verified["rows"] if not o["place"]}

    if verified.get("drive_source") == "estimated":
        client.drive_source = "estimated"
    # ~15 min of driving rarely exceeds ~20 km; drive time is checked precisely below.
    rect = {"low": {"latitude": origin[0] - 0.19, "longitude": origin[1] - 0.23},
            "high": {"latitude": origin[0] + 0.19, "longitude": origin[1] + 0.23}}
    found = {}
    for query, per_city in config.DISCOVERY_QUERIES:
        for q in ([f"{query} in {c}" for c in config.DISCOVERY_CITIES] if per_city else [query]):
            token = None
            for _page in range(config.DISCOVERY_MAX_PAGES):
                res = client.text_search(q, rectangle=rect, page_size=20, page_token=token)
                for p in res.get("places", []):
                    found.setdefault(p["id"], p)
                token = res.get("nextPageToken")
                if not token:
                    break
    print(f"Discovery returned {len(found)} unique places")

    candidates, skipped = [], Counter()
    for p in found.values():
        name = (p.get("displayName") or {}).get("text", "")
        tier, category, nofit = classify(name, p)
        if nofit:
            skipped[nofit] += 1
        elif tier not in ("A", "Review"):
            skipped["not Tier A"] += 1
        elif p.get("businessStatus") != "OPERATIONAL":
            skipped["closed"] += 1
        elif p["id"] in known_ids or digits10(p.get("nationalPhoneNumber")) in known_phones - {""}:
            skipped["already on list"] += 1
        elif clean_name(name) in unverified_names:
            skipped["already on list"] += 1
        elif dnc.match(name=name, place=p):
            skipped["on do-not-contact list"] += 1
        else:
            candidates.append({"place": p, "tier": tier, "category": category, "drive_minutes": None})
    drives = client.drive_minutes(origin, [latlng(c["place"]) for c in candidates])
    new = []
    for c, d in zip(candidates, drives):
        c["drive_minutes"] = d
        if d is not None and d <= config.MAX_DRIVE_MINUTES:
            new.append(c)
        else:
            skipped["over 15 min drive"] += 1
    for c in new:
        c["drive_source"] = client.drive_source
    save_json("discovered.json", new)
    print(f"New Tier A / needs-review organizations: {len(new)}  Skipped: {dict(skipped)}  API calls: {client.calls}")


# --- step: emails --------------------------------------------------------------

def cmd_emails(_args):
    from .emails import find_email
    contacts, _ = assemble()
    emails = load_json("emails.json", default={})
    todo = [c for c in contacts if c["website"] and not c["email"] and c["place_id"] not in emails]
    print(f"Checking {len(todo)} websites for a contact email")
    failed = 0
    for i, c in enumerate(todo, 1):
        email, source = find_email(c["website"])
        if email is None:
            failed += 1   # network failure: leave it out so the next run retries
            continue
        emails[c["place_id"]] = {"email": email, "source": source}
        if i % 25 == 0:
            save_json("emails.json", emails)
            print(f"  {i}/{len(todo)}")
    save_json("emails.json", emails)
    print(f"Emails found: {sum(1 for e in emails.values() if e['email'])}/{len(emails)}")
    if failed:
        print(f"  {failed} websites could not be fetched (network blocked?) — re-run `emails` to retry them")


# --- step: build ---------------------------------------------------------------

def place_fields(place):
    comps = place.get("addressComponents") or []
    city = next((c.get("longText") for c in comps if "locality" in c.get("types", [])), "")
    addr = place.get("formattedAddress", "").replace(", USA", "")
    return {
        "google_name": (place.get("displayName") or {}).get("text", ""),
        "category": (place.get("primaryTypeDisplayName") or {}).get("text", "") or place.get("primaryType", ""),
        "address": addr,
        "city": city,
        "phone": place.get("nationalPhoneNumber", ""),
        "website": place.get("websiteUri", ""),
        "business_status": place.get("businessStatus", ""),
        "place_id": place.get("id", ""),
        "google_maps_url": place.get("googleMapsUri", ""),
    }


def assemble():
    """Merge verified + discovered + emails into (contacts, excluded) row dicts."""
    verified = load_json("verified.json")
    discovered = load_json("discovered.json", default=[])
    emails = load_json("emails.json", default={})
    _, dnc_rows = split_records()
    dnc = DoNotContact(dnc_rows)
    today = date.today().isoformat()

    contacts, excluded = [], []
    seen_ids = {}
    for o in verified["rows"]:
        rec, place = o["record"], o["place"]
        row = {"business": rec["business"], "source": "existing list"}
        if place:
            row.update(place_fields(place))
        else:
            row.update({"phone": pretty_phone(rec["phone"]), "address": rec.get("address", ""),
                        "city": rec.get("city", ""), "website": rec.get("website", ""), "place_id": ""})
        row.update({"email": rec.get("email", ""), "email_source": "existing list" if rec.get("email") else "",
                    "drive_minutes": o["drive_minutes"],
                    "drive_time_source": verified.get("drive_source", "google") if o["drive_minutes"] is not None else "",
                    "verified_by": o["method"] or "not found",
                    "phone_on_file": pretty_phone(rec["phone"]), "verify_note": ""})
        if place and digits10(rec["phone"]) != digits10(place.get("nationalPhoneNumber")):
            row["verify_note"] = "Phone on file differs from Google — use Google number"
        if place and place["id"] in seen_ids:
            row["verify_note"] = (row["verify_note"] + "; " if row["verify_note"] else "") + \
                f"Same Google listing as '{seen_ids[place['id']]}'"
        elif place:
            seen_ids[place["id"]] = rec["business"]
        row.update({k: rec.get(k, "") for k in HISTORY_COLS})
        tier, cat, nofit = classify(rec["business"], place)
        row["tier"], row["tier_category"] = tier, cat if tier in ("A", "Review") else ("Employer (staff on-site)" if tier == "B" else "")
        if tier == "Review":
            note = review_note(rec["business"], place, cat)
            row["verify_note"] = f"{row['verify_note']}; {note}" if row["verify_note"] else note

        hit = dnc.match(phone=rec["phone"], place=place)
        if hit:
            reason, detail = R_DNC, f"Matches do-not-contact record '{hit['business']}'"
        elif not place:
            reason, detail = R_NOTFOUND, "No Google listing matched the name or phone"
        elif place.get("businessStatus") != "OPERATIONAL":
            reason, detail = R_CLOSED, place.get("businessStatus", "")
        elif nofit:
            reason, detail = R_NOFIT, nofit
        elif o["drive_minutes"] is None or o["drive_minutes"] > config.MAX_DRIVE_MINUTES:
            reason, detail = R_FAR, f"{o['drive_minutes']} min" if o["drive_minutes"] else "no drive route"
        else:
            reason, detail = "", ""
        if reason:
            excluded.append({**row, "excluded_reason": reason, "reason_detail": detail})
        else:
            contacts.append(row)

    for d in discovered:
        place = d["place"]
        row = {"business": (place.get("displayName") or {}).get("text", ""), "source": f"new {today}"}
        row.update(place_fields(place))
        tier = d.get("tier", "A")
        note = review_note(row["business"], place, d["category"]) if tier == "Review" else ""
        row.update({"tier": tier, "tier_category": d["category"], "drive_minutes": d["drive_minutes"],
                    "drive_time_source": d.get("drive_source", "google"), "email": "", "email_source": "", "verified_by": "discovery search",
                    "phone_on_file": "", "verify_note": note})
        row.update({k: "" for k in HISTORY_COLS})
        row.update({"status": "prospect", "attempts": "0", "last_outcome": "Never contacted", "events_booked": "0"})
        contacts.append(row)

    for c in contacts:
        e = emails.get(c.get("place_id") or "")
        if e and e["email"] and not c["email"]:
            c["email"], c["email_source"] = e["email"], e["source"]

    for r in dnc_rows:
        row = {"business": r["business"], "source": "do-not-contact list", "phone": pretty_phone(r["phone"]),
               "phone_on_file": pretty_phone(r["phone"]), "email": r.get("email", ""),
               "address": r.get("address", ""), "city": r.get("city", ""), "website": r.get("website", ""),
               "excluded_reason": R_DNC, "reason_detail": r.get("excluded_reason", "")}
        row.update({k: r.get(k, "") for k in HISTORY_COLS})
        excluded.append(row)

    for c in contacts:
        c["priority"] = priority(c["tier_category"]) if c["tier"] in ("A", "Review") else ""
    tier_order = {"A": 0, "Review": 1, "B": 2}
    contacts.sort(key=lambda r: (tier_order.get(r["tier"], 3), r["priority"] or 9,
                                 r.get("drive_minutes") or 99, r["business"].lower()))
    return contacts, excluded


CONTACT_COLS = [
    "tier", "priority", "tier_category", "business", "google_name", "category", "address", "city", "phone",
    "email", "website", "drive_minutes", "drive_time_source", "status", "owner", "last_contacted", "last_outcome",
    "next_step_due", "active_cadence", "attempts", "first_contacted", "last_worked_by",
    "gatekeeper", "decision_maker", "info_sent_to", "last_note", "events_booked",
    "last_event_date", "shares_phone_with_another", "type", "source", "verified_by",
    "verify_note", "phone_on_file", "email_source", "business_status", "place_id", "google_maps_url",
]
EXCLUDED_COLS = ["excluded_reason", "reason_detail", "business", "google_name", "category", "address",
                 "city", "phone", "email", "website", "drive_minutes", "drive_time_source", "business_status"] + \
    [c for c in CONTACT_COLS if c in HISTORY_COLS] + ["source", "verified_by", "phone_on_file", "place_id"]

DAILY_LOG_COLS = ["date", "team_member", "business", "contact_name", "contact_role", "channel",
                  "reached", "outcome", "what_they_said", "next_step", "next_step_due",
                  "handoff_to_caden", "best_days", "notes"]
EVENTS_COLS = ["event_id", "event_date", "start_time", "end_time", "host_business", "event_type",
               "stat_sheet_block", "host_contact", "host_phone", "host_email", "address",
               "presenter", "status", "confirmed_on", "expected_attendance", "actual_attendance",
               "leads_captured", "evals_booked", "notes"]
EVENT_LEADS_COLS = ["lead_id", "event_id", "event_date", "host_business", "first_name", "last_name",
                    "phone", "email", "main_complaint", "interest_level", "first_call_date",
                    "call_outcome", "eval_booked_date", "deposit_taken", "added_to_ltv_cac",
                    "brevo_nurture_added", "notes"]


def write_workbook(contacts, excluded, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title, cols, rows):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") if r.get(c) is not None else "" for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, 1):
            longest = max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows[:300]])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(10, longest + 2), 45)
        return ws

    sheet("Contacts", CONTACT_COLS, contacts)
    sheet("Excluded", EXCLUDED_COLS, excluded)
    log = sheet("Daily Log", DAILY_LOG_COLS, [])
    events = sheet("Events", EVENTS_COLS, [])
    sheet("Event Leads", EVENT_LEADS_COLS, [])

    def dropdown(ws, col_name, cols, options):
        col = ws.cell(1, cols.index(col_name) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}2000")

    dropdown(log, "team_member", DAILY_LOG_COLS, ["Caden", "Dr. Alexa", "Dr. Castillo"])
    dropdown(log, "channel", DAILY_LOG_COLS, ["Phone", "Email", "In person"])
    dropdown(log, "reached", DAILY_LOG_COLS, ["Yes", "No"])
    dropdown(events, "event_type", EVENTS_COLS, ["Massage", "Spinal Screening", "Health Talk"])
    dropdown(events, "stat_sheet_block", EVENTS_COLS, ["Massage", "Spinal Screenings", "Lunch & Learns"])
    dropdown(events, "status", EVENTS_COLS, ["Tentative", "Confirmed", "Completed", "Cancelled"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def summarize(contacts, excluded):
    tiers = Counter(c["tier"] for c in contacts)
    lines = ["CONTACTS", f"  Tier A: {tiers['A']}", f"  Needs review: {tiers['Review']}", f"  Tier B: {tiers['B']}",
             f"  Total:  {len(contacts)}"]
    a_cats = Counter(c["tier_category"] for c in contacts if c["tier"] == "A")
    lines += [f"    A · {k}: {v}" for k, v in a_cats.most_common()]
    src = Counter("new" if c["source"].startswith("new") else "existing" for c in contacts)
    lines += [f"  From existing list: {src['existing']}   New (discovered): {src['new']}",
              f"  With email: {sum(1 for c in contacts if c['email'])}",
              "", "EXCLUDED"]
    lines += [f"  {k}: {v}" for k, v in Counter(e["excluded_reason"] for e in excluded).most_common()]
    lines += [f"  Total: {len(excluded)}"]
    return "\n".join(lines)


def cmd_build(args):
    contacts, excluded = assemble()
    path = config.OUTPUT_DIR / f"castle-outreach-list-{date.today().isoformat()}.xlsx"
    write_workbook(contacts, excluded, path)
    text = summarize(contacts, excluded)
    (config.OUTPUT_DIR / "summary.txt").write_text(text + "\n")
    print(text)
    print(f"\nWrote {path}")


def cmd_all(args):
    cmd_verify(args)
    cmd_discover(args)
    cmd_emails(args)
    cmd_build(args)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("step", choices=["estimate", "verify", "discover", "emails", "build", "all"])
    ap.add_argument("--refresh", action="store_true", help="ignore cached API responses")
    args = ap.parse_args(argv)
    {"estimate": cmd_estimate, "verify": cmd_verify, "discover": cmd_discover,
     "emails": cmd_emails, "build": cmd_build, "all": cmd_all}[args.step](args)


if __name__ == "__main__":
    main()
