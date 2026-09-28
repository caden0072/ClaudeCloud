"""Offline end-to-end test: runs verify → discover → build against a fake Google client."""
import csv
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from castle_prospects import config, gapi, pipeline
from castle_prospects.scoring import classify

ORIGIN = (34.1100, -117.6400)
HDR = ["business", "type", "phone", "email", "website", "address", "city", "drive_minutes", "status",
       "owner", "attempts", "first_contacted", "last_contacted", "last_outcome", "last_worked_by",
       "gatekeeper", "decision_maker", "info_sent_to", "last_note", "events_booked", "last_event_date",
       "active_cadence", "next_step_due", "shares_phone_with_another", "excluded_reason"]


def place(pid, name, phone, types, status="OPERATIONAL", dlat=0.01):
    return {"id": pid, "displayName": {"text": name}, "nationalPhoneNumber": phone, "types": types,
            "primaryType": types[0], "primaryTypeDisplayName": {"text": types[0]},
            "businessStatus": status, "location": {"latitude": ORIGIN[0] + dlat, "longitude": ORIGIN[1]},
            "formattedAddress": f"1 Main St, Upland, CA 91786, USA", "websiteUri": "",
            "addressComponents": [{"longText": "Upland", "types": ["locality"]}]}


PLACES = {
    "Acme Manufacturing": [place("p1", "Acme Manufacturing Inc", "(909) 555-0001", ["manufacturer"])],
    "Iron Temple Gym": [place("p2", "Iron Temple Gym", "(909) 555-0002", ["gym"])],
    "Old Diner": [place("p3", "Old Diner", "(909) 555-0003", ["restaurant"], status="CLOSED_PERMANENTLY")],
    "Far Away Corp": [place("p4", "Far Away Corp", "(909) 555-0004", ["corporate_office"], dlat=0.3)],
    "church in Upland, CA": [place("p10", "Grace Church", "(909) 555-0010", ["church"]),
                              place("p11", "Hard No Church", "(909) 555-0099", ["church"])],
}


class FakeClient:
    def __init__(self, *a, **k):
        self.calls = {}
        self.drive_source = "google"

    def geocode_origin(self, _addr):
        return ORIGIN, "659 E 15th St, Upland"

    def text_search(self, query, **_k):
        return {"places": PLACES.get(query, [])}

    def drive_minutes(self, origin, dests):
        return [round(abs(lat - origin[0]) * 100, 1) for lat, _ in dests]


class PipelineTest(unittest.TestCase):
    def test_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            rows = [("Acme Manufacturing", "+19095550001", "active", "3"),
                    ("Iron Temple Gym", "+19095550002", "prospect", "0"),
                    ("Old Diner", "+19095550003", "prospect", "0"),
                    ("Far Away Corp", "+19095550004", "booked", "2"),
                    ("Ghost LLC", "+19095550005", "prospect", "0"),
                    ("Hard No Church", "+19095550099", "no_event", "1")]
            for fname, subset in (("all.csv", rows), ("dnc.csv", rows[-1:])):
                with open(tmp / fname, "w", newline="") as f:
                    w = csv.DictWriter(f, HDR)
                    w.writeheader()
                    for n, ph, st, att in subset:
                        w.writerow({**{h: "" for h in HDR}, "business": n, "phone": ph, "status": st,
                                    "attempts": att, "last_note": f"note for {n}"})
            with mock.patch.multiple(config, ALL_RECORDS_CSV=tmp / "all.csv", DO_NOT_CONTACT_CSV=tmp / "dnc.csv",
                                     OUTPUT_DIR=tmp / "out", WORK_DIR=tmp / "out" / "work",
                                     DISCOVERY_CITIES=["Upland, CA"], DISCOVERY_QUERIES=[("church", True)]), \
                    mock.patch.object(gapi, "GoogleClient", FakeClient):
                args = mock.Mock(refresh=False)
                pipeline.cmd_verify(args)
                pipeline.cmd_discover(args)
                contacts, excluded = pipeline.assemble()
                pipeline.write_workbook(contacts, excluded, tmp / "out" / "t.xlsx")

            by_name = {c["business"]: c for c in contacts}
            self.assertEqual(by_name["Iron Temple Gym"]["tier"], "A")
            self.assertEqual(by_name["Acme Manufacturing"]["tier"], "B")
            self.assertEqual(by_name["Acme Manufacturing"]["attempts"], "3")          # history kept
            self.assertEqual(by_name["Acme Manufacturing"]["last_note"], "note for Acme Manufacturing")
            self.assertIn("Grace Church", by_name)                                    # discovered
            self.assertNotIn("Hard No Church", by_name)                               # DNC never re-added
            reasons = {e["business"]: e["excluded_reason"] for e in excluded}
            self.assertEqual(reasons["Old Diner"], pipeline.R_CLOSED)
            self.assertEqual(reasons["Far Away Corp"], pipeline.R_FAR)
            self.assertEqual(reasons["Ghost LLC"], pipeline.R_NOTFOUND)
            self.assertEqual(reasons["Hard No Church"], pipeline.R_DNC)

            from openpyxl import load_workbook
            wb = load_workbook(tmp / "out" / "t.xlsx")
            self.assertEqual(wb.sheetnames, ["Contacts", "Excluded", "Daily Log", "Events", "Event Leads"])


class ClassifyTest(unittest.TestCase):
    def test_store_is_not_fitness(self):
        p = {"types": ["home_improvement_store"], "primaryType": "home_improvement_store"}
        self.assertEqual(classify("Lowe's (Upland) fitness studio", p)[0], "B")

    def test_exclusions(self):
        gym = {"types": ["gym", "health"], "primaryType": "gym"}
        self.assertEqual(classify("Fitness Court at Central Park", gym)[0], "")
        self.assertEqual(classify("On Cloud Nine Day Spa", {"types": ["spa", "yoga_studio"], "primaryType": "spa"})[0], "")
        self.assertEqual(classify("New Dawn Sober Living", {"types": ["service"], "primaryType": "service"})[0], "")
        self.assertEqual(classify("SSF Roofers Inc Us", {"types": ["yoga_studio"], "primaryType": "yoga_studio"})[0], "")
        self.assertEqual(classify("Acme Roofing", {"types": ["roofing_contractor"], "primaryType": "roofing_contractor"})[0], "B")
        self.assertEqual(classify("Villa Serena Senior Apartments", {"types": ["apartment_complex"], "primaryType": "apartment_complex"})[2],
                         "Senior community (not pursued)")
        self.assertEqual(classify("Joystar", {"types": ["assisted_living_facility"], "primaryType": "assisted_living_facility"})[0], "")
        self.assertEqual(classify("Bright Haven Hospice", {"types": ["health"], "primaryType": "health"})[0], "")
        self.assertEqual(classify("Retirement Planning Group", {"types": ["financial_planner"], "primaryType": "financial_planner"})[0], "B")
        # big gyms carry a secondary "spa" type for their saunas
        self.assertEqual(classify("Planet Fitness", {"types": ["gym", "spa"], "primaryType": "gym"})[0], "A")

    def test_name_only_match_needs_review_except_civic(self):
        office = {"types": ["service"], "primaryType": "service"}
        self.assertEqual(classify("Sunrise Yoga Collective", office)[0], "Review")
        self.assertEqual(classify("Upland Rotary Club", office)[:2], ("A", "Civic club"))
        self.assertEqual(classify("Grace Church", {"types": ["church"], "primaryType": "church"})[0], "A")

    def test_address_only_is_no_fit(self):
        self.assertTrue(classify("Some Home", {"types": ["premise"], "primaryType": "premise"})[2])


if __name__ == "__main__":
    unittest.main()


class DriveFallbackTest(unittest.TestCase):
    def test_blocked_routes_api_falls_back_to_estimate(self):
        client = gapi.GoogleClient(api_key="x")
        with mock.patch.object(client, "_route_matrix", side_effect=gapi.ServiceBlocked("routes")):
            # ~0.1 degree of latitude north is ~11.1 km straight-line
            mins = client.drive_minutes(ORIGIN, [(ORIGIN[0] + 0.1, ORIGIN[1])])
        self.assertEqual(client.drive_source, "estimated")
        self.assertAlmostEqual(mins[0], 11.1 * config.ROAD_FACTOR / config.AVG_SPEED_KMH * 60, delta=0.3)


class EmailFilterTest(unittest.TestCase):
    def test_skips_legal_and_privacy_addresses(self):
        from castle_prospects.emails import _emails_in
        page = "cpo@24hourfit.com privacy@x.com info@gym.com careers@gym.com"
        self.assertEqual(_emails_in(page), ["info@gym.com"])

